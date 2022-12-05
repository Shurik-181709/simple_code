import math
import openpyxl
from openpyxl.chart import (ScatterChart, Reference, Series)
# Импорт файла с данными
wb = openpyxl.reader.excel.load_workbook(filename="sample_06.xlsx")
wb.active = 0
sheet = wb.active
# Вычисление пиращений коориднат
for j in range(2, 10050):
    H1 = sheet["A" + str(j)].value
    H2 = sheet["A" + str(j + 1)].value
    Z1 = sheet["D" + str(j)].value
    Z2 = sheet["D" + str(j + 1)].value
    A1 = sheet["E" + str(j)].value
    A2 = sheet["E" + str(j + 1)].value
    Dx = (H2 - H1) * (math.sin((Z1 + Z2) / 2)) * (math.cos((A1 + A2) / 2))
    Dy = (H2 - H1) * (math.sin((Z1 + Z2) / 2)) * (math.sin((A1 + A2) / 2))
    Dz = (H2 - H1) * (math.cos((Z1 + Z2) / 2))
    sheet["F" + str(j)].value = Dx
    sheet["G" + str(j)].value = Dy
    sheet["H" + str(j)].value = Dz

# Вычисление координат
for j in range(3, 10050):
    sheet["I2"].value = sheet["F2"].value
    x = sheet["I"+str(j-1)].value + sheet["F"+str(j)].value
    sheet["I"+str(j)].value = x
    sheet["J2"].value = sheet["G2"].value
    y = sheet["J" + str(j - 1)].value + sheet["G" + str(j)].value
    sheet["J" + str(j)].value = y
    sheet["K2"].value = 0 - sheet["H2"].value
    z = sheet["K" + str(j - 1)].value - sheet["H" + str(j)].value
    sheet["K" + str(j)].value = z
# Построение графиков
chart = ScatterChart()
xvalues = Reference(sheet, min_col=9, min_row=2, max_row=10050)
yvalues = Reference(sheet, min_col=10, min_row=2, max_row=10050)
series = Series(yvalues, xvalues)
chart.series.append(series)

sheet.add_chart(chart, "M4")

chart2 = ScatterChart()
xvalues2 = Reference(sheet, min_col=9, min_row=2, max_row=10050)
yvalues2 = Reference(sheet, min_col=11, min_row=2, max_row=10050)
series2 = Series(yvalues2, xvalues2)
chart2.series.append(series2)

sheet.add_chart(chart2, "M18")

chart3 = ScatterChart()
xvalues3 = Reference(sheet, min_col=10, min_row=2, max_row=10050)
yvalues3 = Reference(sheet, min_col=11, min_row=2, max_row=10050)
series3 = Series(yvalues3, xvalues3)
chart3.series.append(series3)

sheet.add_chart(chart3, "M32")

# Сохранение работы в новый файл
wb.save("sample_05.xlsx")



